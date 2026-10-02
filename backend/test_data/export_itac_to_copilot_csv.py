import os
import openpyxl
import csv

def export_itac():
    xlsx_path = "E:/Downloads/ITAC_Database/ITAC_Database_20261001.xlsx"
    out_csv = "E:/Downloads/ITAC_Database/itac_us_doe_22k_copilot.csv"
    
    if not os.path.exists(xlsx_path):
        print(f"File tidak ditemukan di {xlsx_path}")
        return

    print("Membaca sheet ASSESS dari ITAC Database...")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb["ASSESS"]
    
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    idx_id = headers.index("ID")
    idx_fy = headers.index("FY")
    idx_prod = headers.index("PRODLEVEL")
    idx_hours = headers.index("PRODHOURS")
    idx_e2 = headers.index("E2_plant_usage") # Gas usage in MMBTU

    exported = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["timestamp", "production_output", "operating_hours", "fuel_consumption", "fuel_type", "equipment_id"])
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            fy = row[idx_fy]
            prod = row[idx_prod]
            hours = row[idx_hours]
            e2 = row[idx_e2]
            equip_id = row[idx_id] or "Plant-Assessment"

            # Filter valid numerical data
            if prod is not None and hours is not None and e2 is not None:
                try:
                    p_val = round(float(prod), 1)
                    h_val = round(float(hours), 1)
                    f_val = round(float(e2), 1)
                    if p_val > 0 and h_val > 0 and f_val > 0:
                        timestamp_str = f"{fy}-01-01" if fy else "2026-01-01"
                        writer.writerow([timestamp_str, p_val, h_val, f_val, "CNG", equip_id])
                        exported += 1
                except (ValueError, TypeError):
                    continue

    print(f"Berhasil mengekspor {exported:,} baris data audit pabrik ke:")
    print(f" -> {out_csv}")

if __name__ == "__main__":
    export_itac()
